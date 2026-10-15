"""Фирменный бланк исходящего письма и список адресов получателей."""

from __future__ import annotations

import io
import json
from datetime import datetime
from pathlib import Path

from app.config import BACKEND_ROOT

LETTERHEAD_NAME = "Фирменный бланк.docx"
COMPANY_FALLBACK = "НПО «Турбулентность-ДОН»"
SEEDED_RECIPIENTS = ["zhalibin.max@gmail.com"]
_RECIPIENTS_PATH = BACKEND_ROOT / "storage" / "outgoing_recipients.json"

DIRECTION_LABELS = {
    "КоммерческийДиректор": "Коммерческий директор",
    "ПрочиеВнутренние": "Прочие внутренние",
    "ДиректорНПО": "Директор НПО",
    "ТехническийДиректор": "Технический директор",
    "ДиректорПроизводства2": "Директор производства",
    "ФинансовыйДиректор": "Финансовый директор",
    "ЗаместительДиректораПоЭкономическойБезопасности": "Зам. директора по экономической безопасности",
    "ИсполнительныйДиректор": "Исполнительный директор",
}


def direction_label(value: str) -> str:
    text = " ".join(str(value or "").split())
    return DIRECTION_LABELS.get(text, text)


def recipient_emails() -> list[str]:
    stored: list[str] = []
    try:
        raw = json.loads(_RECIPIENTS_PATH.read_text(encoding="utf-8"))
        if isinstance(raw, list):
            stored = [str(item) for item in raw]
    except (OSError, json.JSONDecodeError):
        stored = []
    merged: list[str] = []
    seen: set[str] = set()
    for item in [*SEEDED_RECIPIENTS, *stored]:
        text = " ".join(str(item).split())
        key = text.casefold()
        if "@" not in text or key in seen:
            continue
        seen.add(key)
        merged.append(text)
    return merged


def remember_recipient(email: str) -> None:
    text = " ".join(str(email or "").split())
    if "@" not in text:
        return
    current = recipient_emails()
    if text.casefold() not in {item.casefold() for item in current}:
        current.append(text)
    _RECIPIENTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    _RECIPIENTS_PATH.write_text(json.dumps(current, ensure_ascii=False, indent=1), encoding="utf-8")


def _set_run(run, *, text: str, size: int, bold: bool = False, italic: bool = False, color: str = "") -> None:
    from docx.shared import Pt, RGBColor

    run.text = text
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def build_outgoing_letterhead(
    *,
    organization: str = "",
    theme: str = "",
    body: str = "",
    email: str = "",
    direction: str = "",
    when: datetime | None = None,
) -> bytes:
    """Лист A4: шапка организации, кому, тема и текст письма."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    org = organization.strip() or COMPANY_FALLBACK
    if len(org) == 36 and org.count("-") == 4:
        org = COMPANY_FALLBACK

    head = document.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(head.add_run(), text=org, size=16, bold=True, color="1A3A6B")

    line = document.add_paragraph()
    line.paragraph_format.space_before = Pt(2)
    line.paragraph_format.space_after = Pt(8)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A3A6B")
    borders.append(bottom)
    line._p.get_or_add_pPr().append(borders)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(title.add_run(), text="ИСХОДЯЩЕЕ ПИСЬМО", size=13, bold=True)

    stamp = document.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    moment = when or datetime.now()
    _set_run(stamp.add_run(), text=moment.strftime("%d.%m.%Y"), size=12)

    def meta(label: str, value: str) -> None:
        text = " ".join(value.split())
        if not text:
            return
        row = document.add_paragraph()
        row.paragraph_format.space_after = Pt(2)
        _set_run(row.add_run(), text=f"{label}: ", size=12, bold=True)
        _set_run(row.add_run(), text=text, size=12)

    meta("Кому", email)
    meta("Направление", direction_label(direction))
    meta("Тема", theme)

    document.add_paragraph()
    lines = [item for item in (body or "").splitlines()] or [""]
    for block in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
        paragraph.paragraph_format.space_after = Pt(6)
        _set_run(paragraph.add_run(), text=block, size=12)

    foot = document.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.paragraph_format.space_before = Pt(18)
    _set_run(foot.add_run(), text="Фирменный бланк Оркестратора", size=9, italic=True, color="666666")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def recipients_path() -> Path:
    return _RECIPIENTS_PATH
